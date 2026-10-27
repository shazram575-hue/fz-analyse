from pathlib import Path
import json, sys
import openpyxl
from pypdf import PdfReader
root=Path(sys.argv[1]); out=root/'outputs/MFSD10_程序发现与独立验证_2026-10-26'; raw=root/'原始数据/单细胞/GSE176078'
wb=openpyxl.load_workbook(raw/'41588_2021_911_MOESM4_ESM.xlsx',read_only=True,data_only=True)
info={}
for sh in wb:
    rows=list(sh.iter_rows(values_only=True)); info[sh.title]={'rows':sh.max_row,'cols':sh.max_column,'first_rows':rows[:12]}
    with (out/'metadata'/('supplement_'+sh.title.replace('/','_')+'.tsv')).open('w',encoding='utf-8') as f:
        import csv
        csv.writer(f,delimiter='\t',lineterminator='\n').writerows(rows)
(out/'metadata/Supplement_workbook_inventory.json').write_text(json.dumps(info,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
pdf=PdfReader(raw/'41588_2021_911_MOESM1_ESM.pdf')
(out/'sources/Supplement_methods_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n'+(p.extract_text() or '') for i,p in enumerate(pdf.pages)),encoding='utf-8')
print(json.dumps(info,ensure_ascii=False,default=str,indent=2));print('PDF_PAGES',len(pdf.pages))
